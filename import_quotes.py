import json
import sys
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
SOURCE = Path(sys.argv[1]) if len(sys.argv) > 1 else Path.home() / "Desktop" / "工作簿1.xlsx"
OUTPUT = ROOT / "data" / "quotes.json"
EMPTY_MARKERS = {"", "无", "n/a", "na", "none", "未标注"}


def optional_text(value):
    if value is None or str(value).strip().lower() in EMPTY_MARKERS:
        return ""
    return str(value)


workbook = load_workbook(SOURCE, read_only=True, data_only=True)
sheet = workbook.active
headers = [str(value).strip() if value is not None else "" for value in next(sheet.iter_rows(values_only=True))]
required = ["书名（版本）", "作者", "句子", "章节", "是否入选"]
missing = [name for name in required if name not in headers]
if missing:
    raise ValueError(f"工作簿缺少字段：{', '.join(missing)}")

column = {name: headers.index(name) for name in required}
quotes = []
for source_row, row in enumerate(sheet.iter_rows(values_only=True), start=2):
    selected = optional_text(row[column["是否入选"]]).strip()
    sentence = row[column["句子"]]
    if selected != "是" or sentence is None:
        continue
    quotes.append({
        "id": len(quotes) + 1,
        "book": str(row[column["书名（版本）"]] or ""),
        "author": str(row[column["作者"]] or ""),
        "sentence": str(sentence),
        "chapter": optional_text(row[column["章节"]]),
        "sourceRow": source_row,
    })

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
OUTPUT.write_text(json.dumps(quotes, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"Imported {len(quotes)} selected quotes from {SOURCE.name} to {OUTPUT}")
